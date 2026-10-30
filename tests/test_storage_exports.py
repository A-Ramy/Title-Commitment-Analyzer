import csv
import io
from datetime import date

from openpyxl import load_workbook

from titleagent import demo, pipeline
from titleagent.emailer import summary_html
from titleagent.exports import Checklist, document, tabular
from titleagent.llm import Usage
from titleagent.prompts import REVISED_SEED_LESSONS, SEED_LESSONS


def analysis():
    llm, usage = demo.scripted_llm(), Usage()
    docs = pipeline.Documents(demo.sample_commitment_pdf())
    structure, gaps = pipeline.read_documents(llm, docs, [], usage)
    return pipeline.review(llm, docs, structure, pipeline.DealInfo(closing_date=date(2026, 11, 13)), [], usage, gaps)


def test_revised_seed_lesson_rewritten_in_firestore():
    from tests.conftest import firestore_store

    store = firestore_store()
    old, new = next(iter(REVISED_SEED_LESSONS.items()))
    lid = next(r["id"] for r in store.list_lessons() if r["text"] == new)
    store.set_lesson(lid, text=old)  # a database seeded before the owner changed the rule
    assert old in store.active_lessons()
    store._seed_lessons()  # what every start-up does
    assert store.active_lessons() == SEED_LESSONS


def test_store_round_trip(store):
    assert store.active_lessons() == SEED_LESSONS
    doc = store.save_document(b"%PDF-1.4 x", "c.pdf", "commitment", by="tc@x")
    deal = store.create_deal(by="tc@x", address="1427 Maple Ave", deal_type="Double Close",
                             closing_date=date(2026, 11, 13))
    store.attach_document(doc, deal)
    a = analysis()
    store.save_run(deal, a, doc, None, False, by="tc@x")
    items = store.get_items(deal)
    assert [i.key for i in items] == [i.key for i in a.items]
    store.update_item(deal, items[0].key, by="tc@x", status="Cleared", notes="done")
    assert store.get_items(deal)[0].status == "Cleared"
    assert deal in store.open_items_by_deal()
    fb = store.add_feedback("missing", "Water bill", by="tc@x", deal_id=deal)
    lesson = store.resolve_feedback(fb, True, "Always check final water bills.", by="admin@x")
    assert lesson and store.active_lessons()[-1] == "Always check final water bills."
    assert store.list_feedback("pending") == []
    store.set_setting("tcs", ["Pat"])
    assert store.get_settings()["tcs"] == ["Pat"]
    # what the views read back
    assert store.get_deal(deal)["closing_date"] == date(2026, 11, 13)
    assert store.get_document(doc)["data"] == b"%PDF-1.4 x"
    assert [d["filename"] for d in store.list_documents(deal)] == ["c.pdf"]
    run = store.latest_run(deal)
    assert run["commitment_doc_id"] == doc and len(run["analysis"]["items"]) == len(a.items)
    assert len(store.list_runs(deal)) == 1
    store.update_deal(deal, status="closed")
    assert store.list_deals("open") == [] and [d["id"] for d in store.list_deals()] == [deal]
    extra = items[0].model_copy(update={"key": "missing-1", "title": "Water bill"})
    store.add_item(deal, extra, by="tc@x")
    assert store.get_items(deal)[-1].key == "missing-1"
    lid = store.add_lesson("Manual rule.", by="admin@x")
    store.set_lesson(lid, active=False)
    assert "Manual rule." not in store.active_lessons() and len(store.list_lessons()) == len(SEED_LESSONS) + 2
    ex = store.add_example("Sample", doc, ["1. Payoff"], deal_type="Wholesale")
    store.add_score(ex, 0.9, 9, 10, 11)
    assert store.list_examples()[0]["expert_items"] == ["1. Payoff"] and store.list_scores()[0]["recall"] == 0.9


def test_every_export_format():
    a = analysis()
    cl = Checklist(title="1427 Maple Ave, Columbus, OH", meta=[("Deal type", "Double Close")], items=a.items,
                   coverage_line="B-I 9/9")
    rows = list(csv.reader(io.StringIO(tabular.to_csv(cl, "action").decode("utf-8-sig"))))
    header = rows.index(next(r for r in rows if r and r[0] == "#"))
    body = rows[header + 1:]
    assert len(body) == len([i for i in a.items if i.priority != "fyi"])
    assert body[0][1] == "Closing Blocker"
    wb = load_workbook(io.BytesIO(tabular.to_xlsx(cl)))
    assert wb.active.max_row > len(a.items)
    assert document.to_pdf(cl).startswith(b"%PDF")
    assert document.to_docx(cl)[:2] == b"PK"
    assert cl.filename("pdf") == "1427_Maple_Ave_Columbus_OH_Title_Checklist.pdf"


def test_closing_blockers_only_export():
    a = analysis()
    cl = Checklist(title="1427 Maple Ave, Columbus, OH", meta=[("Deal type", "Double Close")], items=a.items)
    blockers = [i for i in a.items if i.priority == "blocker"]
    assert blockers and len(cl.visible("blockers")) == len(blockers) < len(cl.visible("all"))
    rows = list(csv.reader(io.StringIO(tabular.to_csv(cl, "blockers").decode("utf-8-sig"))))
    assert ["Includes", "Closing Blockers only: Before Closing and Review / FYI items left out"] in rows
    body = rows[rows.index(next(r for r in rows if r and r[0] == "#")) + 1:]
    assert {r[1] for r in body} == {"Closing Blocker"}
    assert cl.filename("pdf", "blockers") == "1427_Maple_Ave_Columbus_OH_Closing_Blockers.pdf"
    assert load_workbook(io.BytesIO(tabular.to_xlsx(cl, "blockers"))).active.title == "Closing Blockers"
    from pypdf import PdfReader

    text = PdfReader(io.BytesIO(document.to_pdf(cl, "blockers"))).pages[0].extract_text()
    assert "Closing Blockers: 1427 Maple Ave" in text and "Before Closing (" not in text
    assert document.to_docx(cl, "blockers")[:2] == b"PK"
    # a clean title: the export says so instead of an empty table
    clean = Checklist(title="Clean St", meta=[], items=[i for i in a.items if i.priority != "blocker"])
    assert "No Closing Blockers on this deal." in PdfReader(io.BytesIO(document.to_pdf(clean, "blockers"))).pages[0] \
        .extract_text()
    assert ["", "No Closing Blockers on this deal."] in list(csv.reader(io.StringIO(
        tabular.to_csv(clean, "blockers").decode("utf-8-sig"))))


def test_daily_email_from_the_store(store):
    """The email built from data exactly as the daily run reads it (closing dates come back as dates)."""
    deal = store.create_deal(by="tc@x", address="9 Sample Rd", deal_type="Double Close",
                             closing_date=date(2026, 11, 20), tc="Pat")
    quiet = store.create_deal(by="tc@x", address="10 Sample Rd", deal_type="Fix & Flip", tc="Sam")
    store.save_run(deal, analysis(), None, None, False, by="tc@x")
    subject, html = summary_html(store.list_deals(status="open"), store.open_items_by_deal(), date(2026, 10, 30),
                                 "https://title.totalchb.com")
    assert "open Closing Blocker" in subject and "9 Sample Rd" in html and "Closing Nov 20 · 21d" in html
    assert "10 Sample Rd" in html and "No Open Blockers" in html and quiet
