"""PDF and Word: a branded, print-ready closing checklist grouped by priority, with a tick box per item."""

from __future__ import annotations

import io
from xml.sax.saxutils import escape

from titleagent.config import ROOT
from titleagent.exports import Checklist, money, source_text
from titleagent.models import PRIORITY_LABEL, due_text

LOGO = ROOT / "assets" / "logo-full.png"
INDIGO, RED, MUTED = "#4D458A", "#BD2239", "#767388"
GROUP_COLOR = {"blocker": RED, "before": "#B45309", "fyi": INDIGO}


def _groups(checklist: Checklist, include: str):
    items = checklist.visible(include)
    for p in ("blocker", "before", "fyi"):
        group = [i for i in items if i.priority == p]
        if group:
            yield p, group


def to_pdf(checklist: Checklist, include: str = "all") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(letter), leftMargin=0.5 * inch, rightMargin=0.5 * inch,
                            topMargin=0.45 * inch, bottomMargin=0.45 * inch,
                            title=f"{checklist.title} - {checklist.heading(include)}")
    body = ParagraphStyle("b", fontName="Helvetica", fontSize=8.5, leading=10.5)
    small = ParagraphStyle("s", parent=body, fontSize=7.5, leading=9, textColor=colors.HexColor(MUTED))
    bold = ParagraphStyle("bb", parent=body, fontName="Helvetica-Bold")
    h1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=16, leading=19, textColor=colors.HexColor(INDIGO))

    story = []
    if LOGO.exists():
        story.append(Image(str(LOGO), width=1.9 * inch, height=1.9 * inch * 274 / 900, hAlign="LEFT"))
        story.append(Spacer(1, 4))
    story.append(Paragraph(f"{checklist.heading(include)}: {escape(checklist.title)}", h1))
    story.append(Spacer(1, 6))
    meta = [*checklist.meta_for(include), ("Coverage", checklist.coverage_line),
            ("Generated", checklist.generated_text)]
    meta = [(k, v) for k, v in meta if v]
    half = (len(meta) + 1) // 2
    left, right = meta[:half], meta[half:]
    rows = []
    for n in range(half):
        lk, lv = left[n]
        rk, rv = right[n] if n < len(right) else ("", "")
        rows.append([Paragraph(f"<b>{escape(lk)}</b>", small), Paragraph(escape(lv), body),
                     Paragraph(f"<b>{escape(rk)}</b>", small), Paragraph(escape(rv), body)])
    t = Table(rows, colWidths=[1.2 * inch, 3.8 * inch, 1.2 * inch, 3.8 * inch])
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
    story += [t, Spacer(1, 10)]

    widths = [0.3, 2.6, 3.25, 1.35, 0.8, 0.95, 0.75]
    if not checklist.visible(include):
        story.append(Paragraph(escape(checklist.empty_note(include)), bold))
    for p, group in _groups(checklist, include):
        color = colors.HexColor(GROUP_COLOR[p])
        story.append(Paragraph(f"{PRIORITY_LABEL[p]} ({len(group)})",
                               ParagraphStyle("g", parent=bold, fontSize=11, leading=14, textColor=color)))
        data = [["", "Item", "Action Needed", "Provided By / Owner", "Due", "Amount", "Status"]]
        for i in group:
            data.append([
                "☐",
                Paragraph(f"<b>{escape(i.title)}</b><br/><font color='{MUTED}' size='7'>{escape(i.category)} · "
                          f"{escape(source_text(i))}</font>", body),
                Paragraph(escape(i.action) + (f"<br/><i>{escape(i.notes)}</i>" if i.notes else ""), body),
                Paragraph(f"{escape(i.provided_by)}<br/><font color='{MUTED}' size='7'>Owner: {escape(i.owner)}"
                          "</font>", body),
                Paragraph(escape(due_text(i)), body), Paragraph(escape(money(i.amount)), body),
                Paragraph(escape(i.status), body),
            ])
        tbl = Table(data, colWidths=[w * inch for w in widths], repeatRows=1)
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(INDIGO)),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 8),
            ("FONTNAME", (0, 1), (0, -1), "DejaVuSans" if _dejavu() else "Helvetica"),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#DDDBEA")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F4FA")]),
        ]))
        story += [tbl, Spacer(1, 10)]
    story.append(Paragraph("Reviewed by: ______________________________ &nbsp;&nbsp; Date: ______________", body))
    doc.build(story)
    return buf.getvalue()


def _dejavu() -> bool:
    """A font with the ☐ tick box glyph, when the system has one."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if "DejaVuSans" in pdfmetrics.getRegisteredFontNames():
        return True
    for path in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/dejavu/DejaVuSans.ttf"):
        try:
            pdfmetrics.registerFont(TTFont("DejaVuSans", path))
            return True
        except Exception:
            continue
    return False


def to_docx(checklist: Checklist, include: str = "all") -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches, Pt, RGBColor

    d = Document()
    sec = d.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.5))
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(9.5)

    if LOGO.exists():
        d.add_picture(str(LOGO), width=Inches(2.0))
    h = d.add_heading(f"{checklist.heading(include)}: {checklist.title}", level=1)
    for r in h.runs:
        r.font.color.rgb = RGBColor(0x4D, 0x45, 0x8A)
    meta = [*checklist.meta_for(include), ("Coverage", checklist.coverage_line),
            ("Generated", checklist.generated_text)]
    t = d.add_table(rows=0, cols=2)
    for k, v in meta:
        if v:
            cells = t.add_row().cells
            cells[0].text, cells[1].text = k, v
            cells[0].paragraphs[0].runs[0].bold = True
    if not checklist.visible(include):
        d.add_paragraph().add_run(checklist.empty_note(include)).bold = True
    for p, group in _groups(checklist, include):
        hh = d.add_heading(f"{PRIORITY_LABEL[p]} ({len(group)})", level=2)
        color = GROUP_COLOR[p].lstrip("#")
        for r in hh.runs:
            r.font.color.rgb = RGBColor.from_string(color.upper())
        tbl = d.add_table(rows=1, cols=7)
        tbl.style = "Light Grid Accent 1"
        for cell, name in zip(tbl.rows[0].cells, ["☐", "Item", "Action Needed", "Provided By / Owner", "Due",
                                                  "Amount", "Status"], strict=True):
            cell.text = name
        for i in group:
            c = tbl.add_row().cells
            c[0].text = "☐"
            c[1].text = f"{i.title}\n{i.category} · {source_text(i)}"
            c[2].text = i.action + (f"\nNote: {i.notes}" if i.notes else "")
            c[3].text = f"{i.provided_by}\nOwner: {i.owner}"
            c[4].text, c[5].text, c[6].text = due_text(i), money(i.amount), i.status
        for row in tbl.rows:
            for w, cell in zip([0.3, 2.6, 3.6, 1.4, 0.8, 0.8, 0.8], row.cells, strict=True):
                cell.width = Inches(w)
    d.add_paragraph()
    d.add_paragraph("Reviewed by: ______________________________    Date: ______________")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
