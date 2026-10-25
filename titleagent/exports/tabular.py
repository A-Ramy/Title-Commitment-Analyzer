"""CSV and Excel."""

from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from titleagent.exports import COLUMNS, Checklist, rows

FILLS = {"Closing Blocker": "FBE3E7", "Before Closing": "FFF1E0", "Review / FYI": "EAE9F4"}
WIDTHS = [5, 16, 22, 44, 60, 8, 22, 12, 12, 11, 22, 30]


def to_csv(checklist: Checklist, include_fyi: bool = True) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([checklist.title])
    for k, v in checklist.meta:
        w.writerow([k, v])
    if checklist.coverage_line:
        w.writerow(["Coverage", checklist.coverage_line])
    w.writerow(["Generated", checklist.generated.strftime("%Y-%m-%d %H:%M")])
    w.writerow([])
    w.writerow(COLUMNS)
    w.writerows(rows(checklist, include_fyi))
    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens it as UTF-8


def to_xlsx(checklist: Checklist, include_fyi: bool = True) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Title Checklist"
    ws.append([checklist.title])
    ws["A1"].font = Font(bold=True, size=14, color="4D458A")
    for k, v in checklist.meta:
        ws.append([k, v])
    if checklist.coverage_line:
        ws.append(["Coverage", checklist.coverage_line])
    ws.append(["Generated", checklist.generated.strftime("%Y-%m-%d %H:%M")])
    ws.append([])
    header_row = ws.max_row + 1
    ws.append(COLUMNS)
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4D458A")
    for r in rows(checklist, include_fyi):
        ws.append(r)
        fill = FILLS.get(r[1])
        if fill:
            ws.cell(row=ws.max_row, column=2).fill = PatternFill("solid", fgColor=fill)
    for i, wdt in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
